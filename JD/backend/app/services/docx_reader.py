"""
docx_reader.py
--------------
Extracts plain text from DOCX files using python-docx.
"""

from docx import Document

from app.utils.helpers import ExtractionError, EmptyFileError


def extract_text_from_docx(file_path: str) -> str:
    """
    Extract and return all text content from a DOCX file, including
    paragraph text and text inside tables.

    Args:
        file_path: Path to the DOCX file on disk.

    Returns:
        The extracted text, stripped of leading/trailing whitespace.

    Raises:
        ExtractionError: if the DOCX cannot be opened or parsed.
        EmptyFileError: if the document contains no extractable text.
    """
    try:
        document = Document(file_path)
    except Exception as exc:
        # Covers corrupted files, files that are not valid .docx
        # (e.g. renamed .doc), etc.
        raise ExtractionError(f"Failed to open DOCX file: {exc}") from exc

    try:
        text_chunks: list[str] = []

        # Regular paragraph text.
        for paragraph in document.paragraphs:
            if paragraph.text and paragraph.text.strip():
                text_chunks.append(paragraph.text.strip())

        # Job descriptions often include tables (e.g. requirements
        # laid out in a grid), so we extract those too.
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text and cell.text.strip():
                        text_chunks.append(cell.text.strip())
    except Exception as exc:
        raise ExtractionError(f"Failed to extract text from DOCX: {exc}") from exc

    full_text = "\n".join(text_chunks).strip()

    if not full_text:
        raise EmptyFileError("No extractable text was found in the DOCX file.")

    return full_text
